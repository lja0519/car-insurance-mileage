"""data/raw_discounts.json → 마일리지_할인_계산기.xlsx

계산기 시트에 할인 전 보험료와 연간 주행거리를 넣으면 보험사별 할인율·할인액·할인 후 금액·순위가
수식으로 계산된다. 할인율은 '할인율표' 시트를 참조하므로 표만 고치면 계산 결과가 따라온다.
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "raw_discounts.json"
OUT = ROOT / "마일리지_할인_계산기.xlsx"
PRESETS = [1000, 2500, 3000, 4500, 5000, 7000, 9000, 11000, 12000, 15000]

STATUS = {"verified": "공식 확인", "secondary": "부분 확인", "unverified": "미확인", "not_offered": "해당없음"}
TABLE_SHEET = "할인율표"

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDE6F7")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
BEST_FILL = PatternFill("solid", fgColor="D6F2DF")
THIN = Side(style="thin", color="C8C8C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, titles):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def basis_date(t):
    if t["valid_from"]:
        return f'{t["valid_from"]}~{t["valid_to"] or ""}'
    return f'~{t["valid_to"]} (시작일 미기재)' if t["valid_to"] else "미기재"


def check_table(ins):
    """수식이 전제하는 조건을 검사한다: 구간이 끊김 없이 이어지고 상한 포함 규칙이 표 안에서 같아야 한다."""
    bands = sorted(ins["primary"]["bands"], key=lambda b: b["distance_min_km"])
    for a, b in zip(bands, bands[1:]):
        if a["distance_max_km"] != b["distance_min_km"]:
            sys.exit(f"{ins['name']}: 구간 공백 {a['distance_max_km']}~{b['distance_min_km']} — 수식 로직 확장 필요")
    rules = {b["max_inclusive"] for b in bands}
    if len(rules) != 1:
        sys.exit(f"{ins['name']}: 상한 포함 규칙이 섞여 있음 {rules} — 수식 로직 확장 필요")
    return bands, rules.pop()


def rate_formula(km, ranges, inclusive, value_range):
    """km 셀 참조에서 할인율(또는 구간 표기)을 찾는 수식. 보간 없이 구간 값만 돌려준다.

    - 상한 '이하'(inclusive=True): km <= 상한인 첫 구간 → COUNTIF(상한,"<"&km)+1
    - 상한 '미만'(inclusive=False): km < 상한인 첫 구간 → COUNTIF(상한,"<="&km)+1
    - 상한 포함 여부 미확인(None): 경계값이면 '경계 불명확', 아니면 구간 값
    - 마지막 상한 초과 → '할인 없음(최대 초과)'. 0% 로 쓰지 않는다.
    """
    max_r, last = ranges["max"], ranges["last"]
    blank = f'OR({km}="",NOT(ISNUMBER({km})),{km}<0)'
    if inclusive is True:
        idx, beyond = f'COUNTIF({max_r},"<"&{km})+1', f"{km}>{last}"
    elif inclusive is False:
        idx, beyond = f'COUNTIF({max_r},"<="&{km})+1', f"{km}>={last}"
    else:
        idx, beyond = f'COUNTIF({max_r},"<"&{km})+1', f"{km}>{last}"
        on_bound = f'COUNTIF({ranges["min"]},{km})+COUNTIF({max_r},{km})>0'
        lookup = f'IF({on_bound},"경계 불명확",INDEX({value_range},{idx}))'
        return f'=IF({blank},"",IF({beyond},"할인 없음(최대 초과)",{lookup}))'
    return f'=IF({blank},"",IF({beyond},"할인 없음(최대 초과)",INDEX({value_range},{idx})))'


def build():
    raw = json.loads(DATA.read_text(encoding="utf-8"))
    insurers = raw["insurers"]
    wb = Workbook()

    # ---------------- 할인율표 (원본 구간) ----------------
    tb = wb.active
    tb.title = TABLE_SHEET
    cols = ["보험사", "구간", "하한(km)", "하한 포함", "상한(km)", "상한 포함", "할인율(%)", "구간 표기",
            "차종", "채널", "할인 방식", "적용 시작일", "적용 종료일", "출처 URL", "확인일", "비교 기준"]
    header(tb, 1, cols)
    row = 2
    ranges = {}
    inc_txt = {True: "포함", False: "미포함", None: "미확인"}
    for ins in insurers:
        tables = ([(ins["primary"], True)] if ins["primary"] else []) + [(t, False) for t in ins["others"]]
        for table, primary in tables:
            src = {s["id"]: s for s in ins["sources"]}
            bands = sorted(table["bands"], key=lambda b: b["distance_min_km"])
            start = row
            for i, b in enumerate(bands, 1):
                tb.append([
                    ins["name"], i, b["distance_min_km"], inc_txt[b["min_inclusive"]], b["distance_max_km"],
                    inc_txt[b["max_inclusive"]], b["discount_rate_percent"], b["label"], table["vehicle_type"],
                    table["channel"], table["discount_method"], table["valid_from"] or "미기재",
                    table["valid_to"] or "", src.get(b["source_id"], {}).get("url", ""), table["verified_at"],
                    "비교 기준" if primary else "별도 표(비교 제외)",
                ])
                row += 1
            if primary:
                check_table(ins)
                end = row - 1
                ranges[ins["id"]] = {
                    "max": f"{TABLE_SHEET}!$E${start}:$E${end}",
                    "rate": f"{TABLE_SHEET}!$G${start}:$G${end}",
                    "label": f"{TABLE_SHEET}!$H${start}:$H${end}",
                    "min": f"{TABLE_SHEET}!$C${start}:$C${end}",
                    "last": f"{TABLE_SHEET}!$E${end}",
                    "inclusive": bands[0]["max_inclusive"] if len({x["max_inclusive"] for x in bands}) == 1 else None,
                }
    autosize(tb, [12, 6, 10, 9, 10, 9, 10, 26, 14, 14, 30, 12, 12, 50, 11, 16])
    tb.freeze_panes = "A2"

    def formula(km, ins, kind):
        r = ranges[ins["id"]]
        return rate_formula(km, r, r["inclusive"], r[kind])

    # ---------------- 계산기 ----------------
    ws = wb.create_sheet("계산기", 0)
    ws["A1"] = "자동차보험 마일리지 특약 할인 계산기"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["A4"] = "할인 전 보험료(원)", "연간 주행거리(km)"
    ws["B3"], ws["B4"] = 1000000, 5000
    for c in ("B3", "B4"):
        ws[c].fill = INPUT_FILL
        ws[c].border = BOX
        ws[c].font = Font(bold=True, size=12)
    ws["B3"].number_format = "#,##0"
    ws["B4"].number_format = "#,##0"
    ws["C3"] = "← 노란 칸에 입력"
    ws["C4"] = "← 노란 칸에 입력 (예상 또는 실제 연간 주행거리)"
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True,
                        error="0 이상의 숫자를 입력하세요")
    ws.add_data_validation(dv)
    dv.add("B3:B4")

    first = 8
    header(ws, first - 1, ["보험사", "할인율(%)", "적용 구간", "할인액(원)", "할인 후 금액(원)", "순위",
                           "할인 방식", "자료 기준일", "검증 상태", "채널"])
    last_row = first + len(insurers) - 1
    rate_col = f"$B${first}:$B${last_row}"
    for i, ins in enumerate(insurers):
        r = first + i
        p = ins["primary"]
        ws.cell(r, 1, ins["name"])
        if ins["id"] in ranges:
            ws.cell(r, 2, formula("$B$4", ins, "rate"))
            ws.cell(r, 3, formula("$B$4", ins, "label"))
        else:
            ws.cell(r, 2, STATUS[ins["status"]])
            ws.cell(r, 3, "할인율표 미확보")
        ws.cell(r, 4, f'=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$3)),ROUND($B$3*B{r}/100,0),IF(B{r}="할인 없음(최대 초과)",0,""))')
        ws.cell(r, 5, f'=IF(ISNUMBER(D{r}),$B$3-D{r},"")')
        ws.cell(r, 6, f'=IF(ISNUMBER(B{r}),RANK(B{r},{rate_col},0),"-")')
        ws.cell(r, 7, p["discount_method"] if p else "-")
        ws.cell(r, 8, basis_date(p) if p else "미확인")
        ws.cell(r, 9, STATUS[ins["status"]])
        ws.cell(r, 10, p["channel"] if p else "-")
        for c in range(1, 11):
            ws.cell(r, c).border = BOX
        ws.cell(r, 4).number_format = "#,##0"
        ws.cell(r, 5).number_format = "#,##0"
        ws.cell(r, 2).alignment = Alignment(horizontal="right")
        ws.cell(r, 6).alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add(f"A{first}:J{last_row}",
                                  FormulaRule(formula=[f'$F{first}=1'], fill=BEST_FILL))

    s = last_row + 2
    ws.cell(s, 1, "최고 할인 보험사").font = BOLD
    ws.cell(s, 2, f'=IF(COUNT({rate_col})=0,"해당 거리 할인 없음",INDEX($A${first}:$A${last_row},MATCH(MAX({rate_col}),{rate_col},0)))')
    ws.cell(s + 1, 1, "최고 할인율(%)").font = BOLD
    ws.cell(s + 1, 2, f'=IF(COUNT({rate_col})=0,"-",MAX({rate_col}))')
    ws.cell(s + 2, 1, "최저 할인 후 금액(원)").font = BOLD
    ws.cell(s + 2, 2, f'=IF(COUNT($E${first}:$E${last_row})=0,"-",MIN($E${first}:$E${last_row}))')
    ws.cell(s + 2, 2).number_format = "#,##0"

    n = s + 4
    ws.cell(n, 1, "유의 사항").font = BOLD
    notes = [
        "할인율은 '할인율표' 시트의 공식 구간표를 그대로 찾아 쓰며, 구간 사이를 보간하지 않습니다.",
        "'할인 없음(최대 초과)': 해당 보험사의 마일리지 할인 최대 거리를 넘은 경우입니다. 할인액 0원으로 계산합니다.",
        "'미확인': 공식 할인율표를 확보하지 못한 보험사입니다. 할인 금액을 계산하지 않습니다(0%로 간주하지 않음).",
        "'경계 불명확': 입력 거리가 구간 경계값인데 약관상 포함 여부가 확인되지 않은 경우입니다.",
        "실제 할인은 보통 마일리지 특약 적용 대상 보험료(일부 담보)에 적용되므로 실제 금액과 다를 수 있습니다.",
        "후할인 방식은 가입 시 할인이 아니라 만기 후 주행거리 확인 뒤 환급됩니다. 할인 방식 열을 확인하세요.",
    ] + raw["meta"]["notes"]
    for i, t in enumerate(notes, 1):
        ws.cell(n + i, 1, f"· {t}")
    autosize(ws, [20, 18, 30, 14, 16, 7, 34, 12, 11, 16])
    ws.freeze_panes = f"A{first}"

    # ---------------- 필수 지점 비교 ----------------
    pv = wb.create_sheet("필수지점_비교", 1)
    pv["A1"] = "필수 비교 지점별 할인율(%) — 계산기와 같은 수식 사용"
    pv["A1"].font = BOLD
    header(pv, 3, ["주행거리(km)"] + [ins["name"] for ins in insurers] + ["최고 보험사"])
    for i, km in enumerate(PRESETS):
        r = 4 + i
        pv.cell(r, 1, km).number_format = "#,##0"
        for j, ins in enumerate(insurers):
            pv.cell(r, 2 + j, formula(f"$A{r}", ins, "rate") if ins["id"] in ranges else STATUS[ins["status"]])
        rng = f"$B{r}:${get_column_letter(1 + len(insurers))}{r}"
        pv.cell(r, 2 + len(insurers),
                f'=IF(COUNT({rng})=0,"-",INDEX($B$3:${get_column_letter(1 + len(insurers))}$3,MATCH(MAX({rng}),{rng},0)))')
    pv.cell(5 + len(PRESETS), 1, "A열 거리를 바꾸면 해당 행이 다시 계산됩니다.")
    autosize(pv, [13] + [12] * len(insurers) + [14])

    # ---------------- Union KM ----------------
    uk = wb.create_sheet("Union_KM")
    union = sorted({b[k] for ins in insurers if ins["primary"] for b in ins["primary"]["bands"]
                    for k in ("distance_min_km", "distance_max_km")})
    header(uk, 1, ["Union KM"] + [ins["name"] for ins in insurers])
    for i, km in enumerate(union):
        r = 2 + i
        uk.cell(r, 1, km).number_format = "#,##0"
        for j, ins in enumerate(insurers):
            uk.cell(r, 2 + j, formula(f"$A{r}", ins, "rate") if ins["id"] in ranges else STATUS[ins["status"]])
    autosize(uk, [11] + [12] * len(insurers))

    # ---------------- 출처 ----------------
    sc = wb.create_sheet("출처")
    header(sc, 1, ["보험사", "검증 상태", "출처 유형", "제목", "URL", "확인일", "인용"])
    for ins in insurers:
        for s_ in ins["sources"] or [{"type": "-", "title": "출처 미확보", "url": "", "accessed_at": "", "quote": ""}]:
            sc.append([ins["name"], STATUS[ins["status"]], s_["type"], s_["title"], s_["url"], s_["accessed_at"],
                       s_.get("quote", "")])
    autosize(sc, [12, 10, 14, 40, 60, 11, 80])

    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    build()
